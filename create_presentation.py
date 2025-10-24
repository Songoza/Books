
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(8.27)
    prs.slide_height = Inches(11.69)

    # Custom colors from the image
    color_background = RGBColor(242, 242, 242)
    color_primary = RGBColor(64, 125, 125)
    color_text = RGBColor(0, 0, 0)

    # Slide layouts
    blank_slide_layout = prs.slide_layouts[6]

    def add_cover_page(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color_primary

        title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        text_frame = title.text_frame
        p = text_frame.paragraphs[0]
        p.text = "WORKBOOK"
        p.font.bold = True
        p.font.size = Pt(44)
        p.font.color.rgb = RGBColor(255, 255, 255)

    def add_content_list_page(slide):
        title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1))
        text_frame = title.text_frame
        p = text_frame.paragraphs[0]
        p.text = "CONTENT LIST"
        p.font.bold = True
        p.font.size = Pt(24)
        p.font.color.rgb = color_primary

    def add_welcome_page(slide):
        title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1))
        text_frame = title.text_frame
        p = text_frame.paragraphs[0]
        p.text = "WELCOME"
        p.font.bold = True
        p.font.size = Pt(24)
        p.font.color.rgb = color_primary

    def add_planner_page(slide, planner_type="WEEKLY"):
        title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1))
        text_frame = title.text_frame
        p = text_frame.paragraphs[0]
        p.text = f"{planner_type} PLANNER"
        p.font.bold = True
        p.font.size = Pt(24)
        p.font.color.rgb = color_primary

    def add_action_list_page(slide):
        title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1))
        text_frame = title.text_frame
        p = text_frame.paragraphs[0]
        p.text = "ACTION LIST"
        p.font.bold = True
        p.font.size = Pt(24)
        p.font.color.rgb = color_primary

    # Create slides
    add_cover_page(prs.slides.add_slide(blank_slide_layout))
    add_content_list_page(prs.slides.add_slide(blank_slide_layout))
    add_welcome_page(prs.slides.add_slide(blank_slide_layout))

    for _ in range(27):
        slide = prs.slides.add_slide(blank_slide_layout)
        # Alternate between different page types for variety
        page_type = _ % 5
        if page_type == 0:
            add_planner_page(slide, "WEEKLY")
        elif page_type == 1:
            add_planner_page(slide, "MONTHLY")
        elif page_type == 2:
            add_action_list_page(slide)
        elif page_type == 3:
            add_welcome_page(slide)
        else:
            add_content_list_page(slide)

    prs.save("workbook_template.pptx")

if __name__ == "__main__":
    create_presentation()
